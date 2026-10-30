"""Buduje public/data/sanok_szpital_jgp.json.gz: wszystkie kody JGP SP ZOZ w Sanoku.

Łączy dwa źródła NFZ po roku, kodzie produktu kontraktowego (zakres = oddział)
i kodzie produktu jednostkowego:
- data/zestawienie_swiadczen_2019_2026.csv — pacjenci i wartość rozliczona,
- data/hospitalizacje_<rok>.csv — liczba hospitalizacji i długość pobytu.

Dodatkowo:
- data/wyniki_finansowe_placowek_*.xlsx — wynik finansowy całego podmiotu,
- data/nfz_umowy_09R_010019.json — umowy z NFZ (scripts/fetch_nfz_agreements.py),
- data/finanse_sp_zoz_sanok.csv — przychody, koszty i wynik netto z kolejnych sprawozdań
  (odczyt ręczny z BizRaport, który publikuje dane z KRS/RDF; źródło i data odczytu w pliku).

"<5" liczymy jako 1 (min) i 4 (max).

Użycie:
    python3 scripts/build_sanok_hospital.py --years 2022,2023,2024,2025
"""

from __future__ import annotations

import argparse
import csv
import glob
import gzip
import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
SETTLEMENTS = DATA_DIR / "zestawienie_swiadczen_2019_2026.csv"
OUTPUT = ROOT / "public" / "data" / "sanok_szpital_jgp.json.gz"

FACILITY = {
    "name": "SP ZOZ w Sanoku",
    "fullName": "Samodzielny Publiczny Zespół Opieki Zdrowotnej w Sanoku",
    "nip": "6871640438",
    "providerCode": "09R/010019",
    "address": "ul. 800-lecia 26, Sanok",
}
LONG_STAY = "6 i więcej dni"

# Oddział szpitalny = grupa zakresów umowy z NFZ (kod VIII z produktu kontraktowego, np. 03.4100.030.02 → 4100).
# Jeden oddział chirurgii w RPWDL (komórka 4500) rozlicza trzy zakresy.
WARDS = [
    {"name": "Choroby wewnętrzne", "contracts": ["4000"]},
    {"name": "Kardiologia", "contracts": ["4100"]},
    {"name": "Neurologia", "contracts": ["4220"]},
    {"name": "Choroby płuc", "contracts": ["4272"]},
    {"name": "Choroby zakaźne", "contracts": ["4340"]},
    {"name": "Pediatria", "contracts": ["4401"]},
    {"name": "Chirurgia ogólna, naczyniowa i urologia", "contracts": ["4500", "4530", "4640"]},
    {"name": "Ortopedia i traumatologia", "contracts": ["4580"]},
    {"name": "Otorynolaryngologia", "contracts": ["4610"]},
]
FIELDS = ["year", "department", "code", "patients", "value", "hospMin", "hospMax", "longStayMin"]

# Grupy finansowania z umów NFZ, w kolejności na wykresie.
AGREEMENT_GROUPS = [
    "Hospitalizacje JGP – ryczałt sieci szpitali",
    "Hospitalizacje JGP poza ryczałtem",
    "SOR",
    "Poradnie i badania",
    "Psychiatria i leczenie uzależnień",
    "Rehabilitacja",
    "Nocna i świąteczna opieka",
    "Dodatki na wynagrodzenia",
    "Inne",
]


def agreement_group(service_type: str, product_code: str) -> str:
    if service_type == "04":
        return "Psychiatria i leczenie uzależnień"
    if service_type == "05":
        return "Rehabilitacja"
    if service_type == "02" or product_code.startswith("02."):
        return "Poradnie i badania"
    if product_code.startswith("03.0000.030"):
        return "Hospitalizacje JGP – ryczałt sieci szpitali"
    if product_code.startswith("03.4"):
        return "Hospitalizacje JGP poza ryczałtem"
    if product_code.startswith("03.3300"):
        return "SOR"
    if product_code.startswith("01.0000.162"):
        return "Nocna i świąteczna opieka"
    if product_code.startswith("00.9999"):
        # Współczynniki korygujące i koszty podwyżek z rozporządzeń (OWU, ustawy o wynagrodzeniach).
        return "Dodatki na wynagrodzenia"
    return "Inne"


def load_finance_history() -> list[dict]:
    """Przychody, koszty i wynik netto z kolejnych sprawozdań finansowych."""
    path = DATA_DIR / "finanse_sp_zoz_sanok.csv"
    if not path.exists():
        return []
    with open(path, encoding="utf-8", newline="") as handle:
        return [
            {
                "year": int(row["rok"]),
                "revenue": int(row["przychody_zl"]),
                "costs": int(row["koszty_zl"]),
                "netResult": int(row["wynik_netto_zl"]),
            }
            for row in csv.DictReader(handle)
        ]


def load_agreements() -> dict | None:
    """Kwoty umów z NFZ w każdym roku, w grupach `AGREEMENT_GROUPS`."""
    path = DATA_DIR / f"nfz_umowy_{FACILITY['providerCode'].replace('/', '_')}.json"
    if not path.exists():
        return None
    totals: dict[int, list[float]] = defaultdict(lambda: [0.0] * len(AGREEMENT_GROUPS))
    for agreement in json.loads(path.read_text(encoding="utf-8"))["agreements"]:
        year, amount = agreement["year"], agreement["amount"]
        plans = agreement["plans"]
        planned = sum(plan["price"] for plan in plans)
        if planned <= 0:
            totals[year][AGREEMENT_GROUPS.index(agreement_group(agreement["serviceType"], ""))] += amount
            continue
        # Pozycje planu dzielą kwotę umowy; skalujemy je do kwoty umowy po aneksach.
        for plan in plans:
            group = agreement_group(agreement["serviceType"], plan["productCode"])
            totals[year][AGREEMENT_GROUPS.index(group)] += amount * plan["price"] / planned
    return {
        "groups": AGREEMENT_GROUPS,
        "years": {str(year): [round(value) for value in values] for year, values in sorted(totals.items())},
    }

CHILDREN = "ŚWIADCZENIA UDZIELANE OSOBOM DO UKOŃCZENIA 18. ROKU ŻYCIA (ŚWIADCZENIA FINANSOWANE ODRĘBNIE W PSZ)"
ENDOPROSTHESIS = ("ŚWIADCZENIA ENDOPROTEZOPLASTYKI STAWU BIODROWEGO LUB KOLANOWEGO "
                  "(ŚWIADCZENIA FINANSOWANE ODRĘBNIE W PSZ)")


def department_name(raw: str) -> str:
    """Krótka nazwa zakresu, np. „Ortopedia i traumatologia narządu ruchu – dzieci do 18 r.ż. (poza PSZ)”."""
    name = re.sub(r"\s*-\s*HOSPITALIZACJA\s*-?\s*", " – ", raw.strip()).strip(" –")
    name = name.replace(CHILDREN, "dzieci do 18 r.ż. (poza PSZ)")
    name = name.replace(ENDOPROSTHESIS, "endoprotezoplastyka biodra lub kolana (poza PSZ)")
    name = name.replace("ORTOPEDIA I TRAUMAT NARZ RUCHU", "ORTOPEDIA I TRAUMATOLOGIA NARZĄDU RUCHU")
    head, *rest = name.split(" – ", 1)
    return head.capitalize() + (f" – {rest[0]}" if rest else "")


def number(raw: str) -> float:
    return float(raw.replace(",", ".")) if raw else 0.0


def contract_unit(contract_code: str) -> str:
    """03.4100.030.02 → 4100."""
    return contract_code.split(".")[1]


def load_finance() -> dict | None:
    """Wynik finansowy podmiotu z najnowszego pliku wyników (arkusz Podmioty_NIP)."""
    from openpyxl import load_workbook

    files = sorted(glob.glob(str(DATA_DIR / "wyniki_finansowe_placowek_*.xlsx")))
    if not files:
        return None
    sheet = load_workbook(files[-1], read_only=True, data_only=True)["Podmioty_NIP"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell) for cell in next(rows)]
    for values in rows:
        row = dict(zip(header, values))
        if str(row.get("NIP")) != FACILITY["nip"]:
            continue
        year_column = next(column for column in header if column.startswith("Wynik netto"))
        return {
            "year": int(re.search(r"\d{4}", year_column).group()),
            "revenue": round(float(row[next(c for c in header if c.startswith("Przychody"))])),
            "netResult": round(float(row[year_column])),
            "source": str(row["Serwis źródłowy"]),
        }
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--years", default="2022,2023,2024,2025", help="lata pliku hospitalizacji")
    args = parser.parse_args()
    hospitalization_years = [year.strip() for year in args.years.split(",") if year.strip()]

    totals: dict[tuple[int, str, str], list[float]] = defaultdict(lambda: [0.0] * 5)
    departments: dict[str, tuple[int, str]] = {}
    products: dict[str, tuple[int, str]] = {}

    with open(SETTLEMENTS, encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            if row["Kod_swiadczeniodawcy"].strip() != FACILITY["providerCode"]:
                continue
            year = int(row["Rok"])
            department = row["Kod_zakresu"].strip()
            product = row["Kod_produktu_jednostkowego"].strip()
            if departments.get(department, (0, ""))[0] <= year:
                departments[department] = (year, row["Nazwa_zakresu"])
            if products.get(product, (0, ""))[0] <= year:
                products[product] = (year, row["Nazwa_produktu_jednostkowego"].strip())
            patients = row["Liczba_pacjentow"].strip()
            entry = totals[(year, department, product)]
            entry[0] += 1 if patients == "<5" else int(patients)
            entry[1] += number(row["Wartosc_jednostek_rozliczonych"])

    for year in hospitalization_years:
        with open(DATA_DIR / f"hospitalizacje_{year}.csv", encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle, delimiter=";"):
                if row["NIP_PODMIOTU"].strip() != FACILITY["nip"]:
                    continue
                raw = row["LICZBA_HOSPITALIZACJI"].strip()
                low, high = (1, 4) if raw == "<5" else (int(raw), int(raw))
                key = (int(row["ROK"]), row["KOD_PRODUKTU_KONTRAKTOWEGO"].strip(),
                       row["KOD_PRODUKTU_JEDNOSTKOWEGO"].strip())
                entry = totals[key]
                entry[2] += low
                entry[3] += high
                if row["PRZEDZIAL_DLUGOSCI_TRWANIA_HOSPITALIZACJI"].strip() == LONG_STAY:
                    entry[4] += low

    department_codes = sorted(departments)

    def ward_of(code: str) -> int | None:
        unit = contract_unit(code)
        return next((index for index, ward in enumerate(WARDS) if unit in ward["contracts"]), None)

    department_index = {code: index for index, code in enumerate(department_codes)}
    codes: dict[str, str] = {}
    for product, (_, full_name) in products.items():
        code, _, name = full_name.partition(" ")
        codes[code] = name.rstrip("* ").strip()

    hospitalization_years_int = {int(year) for year in hospitalization_years}
    facts = []
    for (year, department, product), (patients, value, low, high, long_stay) in sorted(totals.items()):
        if department not in department_index or product not in products:
            print(f"Pominięto wiersz bez rozliczenia: {year} {department} {product}")
            continue
        code = products[product][1].split(" ", 1)[0]
        has_hosp = year in hospitalization_years_int
        facts.append([
            year, department_index[department], code, int(patients), round(value),
            int(low) if has_hosp else None, int(high) if has_hosp else None, int(long_stay) if has_hosp else None,
        ])

    payload = {
        "facility": FACILITY,
        "fields": FIELDS,
        "hospitalizationYears": sorted(hospitalization_years_int),
        # Rok bieżący jest jeszcze rozliczany, więc jego dane są niepełne.
        "partialYears": sorted({fact[0] for fact in facts if fact[0] >= date.today().year}),
        "departments": [
            {"code": code, "name": department_name(departments[code][1]), "ward": ward_of(code)}
            for code in department_codes
        ],
        "wards": [{"name": ward["name"]} for ward in WARDS],
        "finance": load_finance(),
        "agreements": load_agreements(),
        "financeHistory": load_finance_history(),
        "codes": dict(sorted(codes.items())),
        "facts": facts,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(OUTPUT, "wt", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, separators=(",", ":"))
    print(f"Zapisano {OUTPUT.relative_to(ROOT)}: {len(facts)} wierszy, {len(codes)} kodów JGP, "
          f"{len(department_codes)} zakresów.")


if __name__ == "__main__":
    main()
